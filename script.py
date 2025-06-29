import cv2
import numpy as np
import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

# Change file path 
folder_path = r'C:\Users\annab\OneDrive\Desktop\darts_idc_analysis\Imgscans_PRISTINE_sensors'
output_excel = 'rgb_analysis.xlsx'
extension = ['.jpg']

# Store results
results = []

# Loop through images and get mean rgb values
for filename in os.listdir(folder_path):
    if any(filename.lower().endswith(ext) for ext in extension):
        image_path = os.path.join(folder_path, filename)
        img = cv2.imread(image_path)

        if img is None:
            print(f"[!] Failed to load: {filename}")
            continue

        img_rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        R, G, B = cv2.split(img_rgb)

        mean_r = np.mean(R)
        mean_g = np.mean(G)
        mean_b = np.mean(B)

        results.append({
            'Image': filename,
            'Mean_R': round(mean_r, 2),
            'Mean_G': round(mean_g, 2),
            'Mean_B': round(mean_b, 2)
        })

# Export DataFrame to Excel, and add bar chart
df = pd.DataFrame(results)
df.to_excel(output_excel, index=False)
wb = load_workbook(output_excel)
ws = wb.active

# Define the data range for the chart
data = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=1 + len(df))
cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(df))

chart = BarChart()
chart.title = "Mean RGB Values per Image"
chart.x_axis.title = "Image"
chart.y_axis.title = "Mean Intensity"

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 10
chart.width = 20

ws.add_chart(chart, f"A{len(df) + 5}")
wb.save(output_excel)

